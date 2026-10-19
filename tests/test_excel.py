"""excel 模块的行为规格。

测试只穿过约定的 seam：`perf_compare.excel.write_excel`；
读取产出文件用 openpyxl，因为 xlsx 本身就是对外契约。

工作簿布局（需求方修正）：
- 「汇总」        —— 套级唯一值的对比（不是劣化项计数）
- 「子项劣化汇总」  —— 跨套列出被判定劣化的子项
- 每个测试套一个 sheet —— 该套全部子项的明细

列结构：测试项与子项**合并成一列**，用 `-` 连接（需求方 2026-10-05）：
`randread_bw_blocksize` + `blocksize-1024k` → `randread_bw_blocksize-blocksize-1024k`。
"""
import shutil
from pathlib import Path

import openpyxl
import pytest

from perf_compare.compare import compare
from perf_compare.excel import write_excel
from perf_compare.resultdir import load
from perf_compare.spark_items import SPARK_ITEMS

FIXTURE = Path(__file__).resolve().parents[1] / "测试结果模版"

明细表头 = ["测试子项", "结果A", "结果B", "差距百分比", "是否劣化"]


def _扰动(path, item, sub, factor):
    lines = path.read_text(encoding="utf-8").splitlines()
    for index, line in enumerate(lines):
        parts = line.split()
        if parts and parts[0] == item:
            position = parts[1:].index(sub)
            average = lines[index + 1].split()
            average[position + 1] = repr(float(average[position + 1]) * factor)
            lines[index + 1] = " ".join(average)
            break
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _两侧样例(tmp_path, 文件名="fio.txt"):
    for side in ("baseline", "candidate"):
        target = tmp_path / side / "基础性能测试"
        target.mkdir(parents=True)
        shutil.copy(str(FIXTURE / "基础性能测试" / 文件名), str(target / 文件名))


def _写出(tmp_path):
    comparison = compare(
        load(str(tmp_path / "baseline")),
        load(str(tmp_path / "candidate")),
        threshold_pct=5.0,
    )
    out_path = tmp_path / "对比表格.xlsx"
    write_excel(comparison, str(out_path))
    return openpyxl.load_workbook(str(out_path))


def _明细行(workbook, sheet):
    return [
        [cell.value for cell in row] for row in workbook[sheet].iter_rows(min_row=2)
    ]


def test_写出对比表格_表头名称可配置(tmp_path):
    """需求方要求：表头默认是「结果A／结果B」，但名称要能指定（如系统名）。"""
    _两侧样例(tmp_path)
    comparison = compare(
        load(str(tmp_path / "baseline")),
        load(str(tmp_path / "candidate")),
        threshold_pct=5.0,
    )
    out_path = tmp_path / "对比表格.xlsx"

    write_excel(comparison, str(out_path), labels=("openEuler-24.03", "openEuler-24.03-SP1"))

    workbook = openpyxl.load_workbook(str(out_path))
    assert [cell.value for cell in workbook["汇总"][1]][:3] == [
        "测试套",
        "openEuler-24.03",
        "openEuler-24.03-SP1",
    ]
    assert [cell.value for cell in workbook["fio"][1]][:3] == [
        "测试子项",
        "openEuler-24.03",
        "openEuler-24.03-SP1",
    ]
    assert [cell.value for cell in workbook["子项劣化汇总"][1]][:4] == [
        "测试套",
        "测试子项",
        "openEuler-24.03",
        "openEuler-24.03-SP1",
    ]


def test_写出对比表格_表头名称的默认值仍是结果A与结果B(tmp_path):
    """不配置名称时必须保持原样，避免影响既有用法。"""
    _两侧样例(tmp_path)
    comparison = compare(
        load(str(tmp_path / "baseline")),
        load(str(tmp_path / "candidate")),
        threshold_pct=5.0,
    )
    out_path = tmp_path / "对比表格.xlsx"

    write_excel(comparison, str(out_path))

    workbook = openpyxl.load_workbook(str(out_path))
    assert [cell.value for cell in workbook["汇总"][1]][:3] == ["测试套", "结果A", "结果B"]


def test_写出对比表格_劣化幅度用分档底色且不用条件格式(tmp_path):
    """兼容性优先（需求方 2026-10-16）：差距列按劣化幅度用**纯静态填充**分档，
    不使用条件格式（数据条在 WPS 与 Excel 上表现可能不一致）。

    算例：同一测试项下三个子项分别劣化 60% / 30% / 10%。
    """
    _两侧样例(tmp_path)
    fio = tmp_path / "candidate" / "基础性能测试" / "fio.txt"
    _扰动(fio, "read_iops_blocksize", "blocksize-4k", 0.40)    # -60%
    _扰动(fio, "read_iops_blocksize", "blocksize-16k", 0.70)   # -30%
    _扰动(fio, "read_iops_blocksize", "blocksize-32k", 0.90)   # -10%
    comparison = compare(
        load(str(tmp_path / "baseline")),
        load(str(tmp_path / "candidate")),
        threshold_pct=5.0,
    )
    out_path = tmp_path / "对比表格.xlsx"

    write_excel(comparison, str(out_path))

    sheet = openpyxl.load_workbook(str(out_path))["fio"]
    fills = {}
    for row in sheet.iter_rows(min_row=2):
        if row[0].value.startswith("read_iops_blocksize-"):
            fills[row[0].value.split("-", 1)[1]] = str(row[3].fill.fgColor.rgb)

    assert fills["blocksize-4k"].endswith("C00000")     # ≥50% 深红
    assert fills["blocksize-16k"].endswith("FF7C80")    # ≥20% 中红
    assert fills["blocksize-32k"].endswith("FFC7CE")    # 超阈值 浅红

    # 全工作簿不使用条件格式
    for name in sheet.parent.sheetnames:
        assert not list(sheet.parent[name].conditional_formatting)


def test_写出对比表格_说明首页与表格美化(tmp_path):
    """必做美化项（需求方 2026-10-16）：

    说明首页、表头样式、冻结首行、自动筛选、列宽、劣化标红、差距带符号百分号、
    劣化汇总按严重度倒序、汇总表类别分组与套名超链接。
    """
    _两侧样例(tmp_path)
    fio = tmp_path / "candidate" / "基础性能测试" / "fio.txt"
    _扰动(fio, "read_iops_blocksize", "blocksize-4k", 0.80)    # 劣化 20%
    _扰动(fio, "read_iops_blocksize", "blocksize-16k", 0.85)   # 劣化 15%
    comparison = compare(
        load(str(tmp_path / "baseline")),
        load(str(tmp_path / "candidate")),
        threshold_pct=5.0,
    )
    out_path = tmp_path / "对比表格.xlsx"

    write_excel(
        comparison,
        str(out_path),
        labels=("基线版本X", "对比版本Y"),
        info={
            "baseline_path": "/data/base",
            "candidate_path": "/data/cand",
            "threshold_pct": 5.0,
        },
    )
    workbook = openpyxl.load_workbook(str(out_path))

    # 说明首页：两侧名称、路径、阈值都在
    assert workbook.sheetnames[0] == "说明"
    notes_text = " ".join(
        str(cell.value)
        for row in workbook["说明"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "基线版本X" in notes_text
    assert "/data/base" in notes_text
    assert "5.0%" in notes_text

    # 表头样式 + 冻结 + 筛选 + 列宽
    summary = workbook["汇总"]
    header = summary.cell(row=1, column=1)
    assert header.font.bold is True
    assert str(header.fill.fgColor.rgb).endswith("1F4E79")
    assert summary.freeze_panes == "A2"
    assert summary.auto_filter.ref is not None
    assert summary.column_dimensions["A"].width >= 14

    # 汇总表按类别分组 + 套名超链接
    summary_labels = [row[0] for row in _明细行(workbook, "汇总")]
    assert summary_labels[0] == "基础性能测试"
    assert "fio" in summary_labels
    fio_cells = [
        cell for row in summary.iter_rows() for cell in row if cell.value == "fio"
    ]
    assert fio_cells[0].hyperlink is not None

    # 明细表：劣化行标红、差距列带符号百分号
    detail = workbook["fio"]
    detail_rows = [row for row in detail.iter_rows(min_row=2)]
    degraded = [row for row in detail_rows if row[4].value == "是"]
    assert len(degraded) == 2
    assert str(degraded[0][0].fill.fgColor.rgb).endswith("FFC7CE")
    assert degraded[0][3].number_format == '+0.00"%";-0.00"%";0.00"%"'

    # 劣化汇总按严重度倒序：-20% 在 -15% 前面
    degraded_sheet = workbook["子项劣化汇总"]
    percentages = [row[4].value for row in degraded_sheet.iter_rows(min_row=2)]
    assert percentages == [-20.0, -15.0]


def test_写出对比表格_spark与spark_docker用宽表明细(tmp_path):
    """spark 明细：一行一个测试项，各轮与平均值**横向并列**（需求方 2026-10-12）。

    列结构 = 测试子项 | 结果A 的各次测量 + 平均 | 结果B 同 | 差距百分比 | 是否劣化；
    行顺序按源文件顺序（`spark_items.py`），且**不含**「汇总」项。
    """
    comparison = compare(load(str(FIXTURE)), load(str(FIXTURE)), threshold_pct=5.0)
    out_path = tmp_path / "对比表格.xlsx"

    write_excel(comparison, str(out_path))

    workbook = openpyxl.load_workbook(str(out_path))

    sheet = workbook["spark"]
    assert [cell.value for cell in sheet[1]] == [
        "测试子项",
        "结果A-轮1", "结果A-轮2", "结果A-轮3", "结果A-avg",
        "结果B-轮1", "结果B-轮2", "结果B-轮3", "结果B-avg",
        "差距百分比",
        "是否劣化",
    ]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
    assert [row[0] for row in rows] == list(SPARK_ITEMS)
    first = rows[0]
    assert first[1] == 33.498
    assert first[2] == 31.663
    assert first[3] == 30.929
    assert first[4] == pytest.approx(32.03)
    assert first[10] == "否"

    docker_sheet = workbook["spark_docker"]
    assert [cell.value for cell in docker_sheet[1]] == [
        "测试子项",
        "结果A-docker0", "结果A-docker1", "结果A-docker2", "结果A-avg",
        "结果B-docker0", "结果B-docker1", "结果B-docker2", "结果B-avg",
        "差距百分比",
        "是否劣化",
    ]
    docker_rows = [[cell.value for cell in row] for row in docker_sheet.iter_rows(min_row=2)]
    assert docker_rows[0][0] == "native_q1"
    assert docker_rows[0][1] == 24.92
    assert docker_rows[0][2] == 14.981
    assert docker_rows[0][3] == 23.628
    assert docker_rows[0][4] == pytest.approx(21.176333333333336)

    # 格式必须落在「差距百分比」列上，不能错加到平均值列（曾因列号算错而错位：
    # 平均列被显示成 `+40.04%`，真正的差距列反而没有 %）。
    percent_format = '+0.00"%";-0.00"%";0.00"%"'
    for wide, avg_columns in ((sheet, (5, 9)), (docker_sheet, (5, 9))):
        assert wide.cell(row=2, column=10).value is not None
        assert wide.cell(row=2, column=10).number_format == percent_format
        for column in avg_columns:
            assert wide.cell(row=2, column=column).number_format == "General"


def test_写出对比表格_输出目录不存在时自动创建(tmp_path):
    """输出目录不存在就直接失败，是个不必要的坑 —— 应当自动建目录。"""
    _两侧样例(tmp_path)
    comparison = compare(
        load(str(tmp_path / "baseline")),
        load(str(tmp_path / "candidate")),
        threshold_pct=5.0,
    )
    out_path = tmp_path / "还不存在" / "更深一层" / "对比表格.xlsx"

    write_excel(comparison, str(out_path))

    assert out_path.is_file()


def test_写出对比表格_flink_wordcount的单位换算写成公式(tmp_path):
    """需求方要求 ÷1000 在 Excel 里**以公式呈现、保留原值**（决策 D19）。

    单元格里应当是 `=116445015/1000`，而不是算好的 116445.015。
    """
    comparison = compare(load(str(FIXTURE)), load(str(FIXTURE)), threshold_pct=5.0)
    out_path = tmp_path / "对比表格.xlsx"

    write_excel(comparison, str(out_path))

    rows = list(
        openpyxl.load_workbook(str(out_path))["汇总"].iter_rows(values_only=True)
    )
    flink_row = [row for row in rows if row[0] == "flink_wordcount"][0]
    assert flink_row[1] == "=116445015/1000"
    assert flink_row[2] == "=116445015/1000"


def test_写出对比表格_汇总表用分组标签作为行名(tmp_path):
    """unixbench 在汇总里占单核／多核两行，行名必须能区分，不能都叫 unixbench。"""
    comparison = compare(load(str(FIXTURE)), load(str(FIXTURE)), threshold_pct=5.0)
    out_path = tmp_path / "对比表格.xlsx"

    write_excel(comparison, str(out_path))

    labels = [
        [cell.value for cell in row][0]
        for row in openpyxl.load_workbook(str(out_path))["汇总"].iter_rows(min_row=2)
    ]
    assert "unixbench（单核）" in labels
    assert "unixbench（多核）" in labels
    assert "unixbench" not in labels


def test_写出对比表格_明细表把测试项与子项合并成一列(tmp_path):
    _两侧样例(tmp_path)

    workbook = _写出(tmp_path)

    sheet = workbook["fio"]
    assert [cell.value for cell in sheet[1]] == 明细表头
    labels = [row[0] for row in _明细行(workbook, "fio")]
    assert "randread_bw_blocksize-blocksize-1024k" in labels
    assert "read_iops_blocksize-blocksize-4k" in labels


def test_写出对比表格_首表是套级汇总(tmp_path):
    _两侧样例(tmp_path)
    _扰动(tmp_path / "candidate" / "基础性能测试" / "fio.txt", "read_iops_blocksize", "blocksize-4k", 0.8)

    workbook = _写出(tmp_path)

    # 需求变更（2026-10-16）：新增「说明」首页，首个 sheet 不再是「汇总」。
    assert workbook.sheetnames[0] == "说明"
    assert "fio" in workbook.sheetnames

    summary = workbook["汇总"]
    # 表头是 A–E 五列；汇总行按类别分组，故首行数据是类别标题。
    assert [cell.value for cell in summary[1]][:5] == [
        "测试套",
        "结果A",
        "结果B",
        "差距百分比",
        "是否劣化",
    ]
    assert [row[0] for row in _明细行(workbook, "汇总")] == ["基础性能测试", "fio"]


def test_写出对比表格_子项劣化汇总表只列被判为劣化的子项(tmp_path):
    _两侧样例(tmp_path)
    _扰动(tmp_path / "candidate" / "基础性能测试" / "fio.txt", "read_iops_blocksize", "blocksize-4k", 0.8)

    workbook = _写出(tmp_path)

    assert "子项劣化汇总" in workbook.sheetnames
    sheet = workbook["子项劣化汇总"]
    assert [cell.value for cell in sheet[1]] == [
        "测试套",
        "测试子项",
        "结果A",
        "结果B",
        "差距百分比",
    ]

    rows = _明细行(workbook, "子项劣化汇总")
    assert len(rows) == 1
    assert rows[0][0] == "fio"
    assert rows[0][1] == "read_iops_blocksize-blocksize-4k"
    assert rows[0][4] == pytest.approx(-20.0, rel=1e-9)


def test_写出对比表格_差距百分比保留两位小数(tmp_path):
    """0.92 倍就是精确的 -8%，但浮点会算成 -7.999999999999999，交付物里必须写成 -8.0。"""
    _两侧样例(tmp_path, "stream.txt")
    _扰动(tmp_path / "candidate" / "基础性能测试" / "stream.txt", "384_Core", "Copy", 0.92)

    workbook = _写出(tmp_path)

    rows = _明细行(workbook, "stream")
    hit = [row for row in rows if row[0] == "384_Core-Copy"]
    assert len(hit) == 1
    assert hit[0][3] == -8.0


def test_写出对比表格_汇总表的套级唯一值取子项几何平均(tmp_path):
    """合成样例便于独立验算：套内两个子项 100 与 400，几何平均为 √(100×400)=200；
    对比侧改成 100 与 100，几何平均为 100；故套级差距 = (100−200)/200 = −50%。
    """
    for side in ("baseline", "candidate"):
        (tmp_path / side / "基础性能测试").mkdir(parents=True)
    (tmp_path / "baseline" / "基础性能测试" / "stream.txt").write_text(
        "1_Core a b\naverage 100.0 400.0\nvariance 0.0 0.0\n", encoding="utf-8"
    )
    (tmp_path / "candidate" / "基础性能测试" / "stream.txt").write_text(
        "1_Core a b\naverage 100.0 100.0\nvariance 0.0 0.0\n", encoding="utf-8"
    )

    workbook = _写出(tmp_path)

    summary = workbook["汇总"]
    rows = [row for row in summary.iter_rows(min_row=2, values_only=True)]
    stream_row = [row for row in rows if row[0] == "stream"][0]
    # 几何平均在对数域计算（避免大数值连乘溢出），故有极小浮点尾差。
    assert stream_row[1] == pytest.approx(200.0, rel=1e-12)
    assert stream_row[2] == pytest.approx(100.0, rel=1e-12)
    assert stream_row[3] == pytest.approx(-50.0, rel=1e-9)
    assert stream_row[4] == "是"
    # 需求变更（2026-10-16）：口径说明从「汇总」的 F1 移到「说明」页。
    notes_text = " ".join(
        str(cell.value)
        for row in workbook["说明"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "口径按测试套声明" in notes_text


def test_写出对比表格_套级几何平均不会因数值过大而溢出(tmp_path):
    """fio 有 96 个子项、量级约 1e5，朴素连乘会溢出成 inf，而 Excel 存不下 inf（读回来是空值）。
    算例：两个 1e300，几何平均应为 1e300。
    """
    for side in ("baseline", "candidate"):
        (tmp_path / side / "基础性能测试").mkdir(parents=True)
        (tmp_path / side / "基础性能测试" / "stream.txt").write_text(
            "1_Core a b\naverage 1e300 1e300\nvariance 0.0 0.0\n", encoding="utf-8"
        )

    summary = _写出(tmp_path)["汇总"]
    rows = list(summary.iter_rows(min_row=2, values_only=True))
    stream_row = [row for row in rows if row[0] == "stream"][0]
    assert stream_row[1] == pytest.approx(1e300, rel=1e-9)
    assert stream_row[2] == pytest.approx(1e300, rel=1e-9)


def test_写出对比表格_百分比无定义的子项显示为无法比较(tmp_path):
    """基线为 0 时留空白会被误读成漏数据，所以显式写出「无法比较」。"""
    _两侧样例(tmp_path, "lmbench3.txt")

    rows = _明细行(_写出(tmp_path), "lmbench3")
    hit = [row for row in rows if row[0] == "Local_latencies-RPC/UDP"]
    assert len(hit) == 1
    assert hit[0][3] == "无法比较"
    assert hit[0][4] == "否"


def test_写出对比表格_单边缺失的测量项显式标注而不是留空白(tmp_path):
    """核数不同会让项名对不上（384_Core vs 256_Core），缺失的一侧要写明「缺失」。"""
    for side, item in (("baseline", "384_Core"), ("candidate", "256_Core")):
        target = tmp_path / side / "基础性能测试"
        target.mkdir(parents=True)
        (target / "stream.txt").write_text(
            "{0} Copy\naverage 100.0\nvariance 0.0\n".format(item), encoding="utf-8"
        )

    rows = _明细行(_写出(tmp_path), "stream")
    by_label = {row[0]: row for row in rows}
    assert by_label["384_Core-Copy"][1] == 100.0
    assert by_label["384_Core-Copy"][2] == "缺失"
    assert by_label["256_Core-Copy"][1] == "缺失"
    assert by_label["256_Core-Copy"][2] == 100.0
