"""cli 模块的行为规格（端到端）。

测试只穿过约定的 seam：`perf_compare.cli.main`。
"""
import json
import shutil
from pathlib import Path

import openpyxl

from perf_compare.cli import main

FIXTURE = Path(__file__).resolve().parents[1] / "测试结果模版"


def test_命令行_读配置并写出对比表格(tmp_path):
    """一条命令跑通：配置 → 解析 → 比对 → 写出 xlsx。"""
    baseline = tmp_path / "baseline"
    candidate = tmp_path / "candidate"
    (baseline / "基础性能测试").mkdir(parents=True)
    (candidate / "基础性能测试").mkdir(parents=True)

    source = (FIXTURE / "基础性能测试" / "fio.txt").read_text(encoding="utf-8")
    (baseline / "基础性能测试" / "fio.txt").write_text(source, encoding="utf-8")

    lines = source.splitlines()
    fields = lines[1].split()
    fields[1] = repr(float(fields[1]) * 0.8)
    lines[1] = " ".join(fields)
    (candidate / "基础性能测试" / "fio.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")

    config_path = tmp_path / "compare.json"
    config_path.write_text(
        json.dumps(
            {
                "baseline": {"path": str(baseline), "kernel_version": "6.6.0-72.3.0"},
                "candidates": [{"path": str(candidate), "kernel_version": "6.6.0-72.5.0"}],
                "threshold_pct": 5.0,
            }
        ),
        encoding="utf-8",
    )
    excel_path = tmp_path / "对比表格.xlsx"

    exit_code = main(["--config", str(config_path), "--excel", str(excel_path)])

    assert exit_code == 0
    assert excel_path.is_file()


def test_命令行_出错时把堆栈写进日志并返回非零(tmp_path):
    """报错要落到本地日志文件里，便于事后排查，而不是只在终端一闪而过。"""
    log_dir = tmp_path / "logs"
    missing_config = tmp_path / "不存在的配置.json"

    exit_code = main(
        [
            "--config",
            str(missing_config),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(log_dir),
        ]
    )

    assert exit_code != 0
    log_text = (log_dir / "perf-compare.log").read_text(encoding="utf-8")
    assert "不存在的配置.json" in log_text
    assert "Traceback" in log_text


def _写配置(tmp_path, baseline, candidate):
    config_path = tmp_path / "compare.json"
    config_path.write_text(
        json.dumps(
            {
                "baseline": {"path": str(baseline)},
                "candidates": [{"path": str(candidate)}],
                "threshold_pct": 5.0,
            }
        ),
        encoding="utf-8",
    )
    return config_path


def test_命令行_一个套都没读到时报错而不是静默产出空表(tmp_path, capsys):
    """最危险的失败方式是「退出码 0 + 空表」：路径填错时用户以为跑成功了。"""
    empty = tmp_path / "空目录"
    empty.mkdir()
    config_path = _写配置(tmp_path, empty, empty)

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    assert exit_code != 0
    assert "结果目录" in capsys.readouterr().err


def test_命令行_文件在但解析不出数据时单独报告(tmp_path, capsys):
    """文件存在、但内容不符合预期格式时，不能笼统说「未找到」——
    否则用户无法判断是文件没生成、名字不对，还是格式变了。"""
    result_dir = tmp_path / "结果"
    result_dir.mkdir()
    shutil.copy(str(FIXTURE / "基础性能测试" / "fio.txt"), str(result_dir / "fio.txt"))
    (result_dir / "mysql_tpcc.txt").write_text("格式不对的内容\n", encoding="utf-8")
    config_path = _写配置(tmp_path, result_dir, result_dir)

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "mysql_tpcc.txt" in output
    assert "没解析出数据" in output


def test_命令行_表头名称来自配置(tmp_path):
    """两侧名称支持在配置里指定（需求方 2026-10-12）。

    注意：断言用的名称**不能出现在测试函数名里** —— pytest 会用函数名创建 tmp_path，
    而 CLI 会把路径打印出来，否则子串断言会被目录名意外满足（假通过）。
    """
    for side in ("b", "c"):
        target = tmp_path / side / "基础性能测试"
        target.mkdir(parents=True)
        shutil.copy(str(FIXTURE / "基础性能测试" / "fio.txt"), str(target / "fio.txt"))

    config_path = tmp_path / "compare.json"
    config_path.write_text(
        json.dumps(
            {
                "baseline": {"path": str(tmp_path / "b"), "name": "openEuler-24.03-64K"},
                "candidates": [
                    {"path": str(tmp_path / "c"), "name": "openEuler-24.03-SP1-64K"}
                ],
                "threshold_pct": 5.0,
            }
        ),
        encoding="utf-8",
    )
    excel_path = tmp_path / "对比表格.xlsx"

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(excel_path),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    assert exit_code == 0
    workbook = openpyxl.load_workbook(str(excel_path))
    assert [cell.value for cell in workbook["汇总"][1]][:3] == [
        "测试套",
        "openEuler-24.03-64K",
        "openEuler-24.03-SP1-64K",
    ]
    assert [cell.value for cell in workbook["fio"][1]][:3] == [
        "测试子项",
        "openEuler-24.03-64K",
        "openEuler-24.03-SP1-64K",
    ]


def test_命令行_无效结果给出告警(tmp_path, capsys):
    """带无效标记的结果文件要提示 —— 其汇总值是估计值、不可作为可报告结果
    （需求方 2026-10-16 确认要告警）。"""
    result_dir = tmp_path / "结果"
    result_dir.mkdir()
    shutil.copy(
        str(FIXTURE / "基础性能测试" / "cpu2017_int_single.txt"),
        str(result_dir / "cpu2017_int_single.txt"),
    )
    config_path = _写配置(tmp_path, result_dir, result_dir)

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    warned = [line for line in output.splitlines() if line.startswith("    ⚠")]
    assert warned, "实际输出为：\n" + output
    # 断言按内容匹配而不是取第一条：该用例两侧同目录，会先打印「同一个目录」告警。
    invalid_lines = [line for line in warned if "INVALID RUN" in line]
    assert invalid_lines, "实际输出为：\n" + output
    assert "speccpu2017_int_single" in invalid_lines[0]


def test_命令行_两侧指向同一个目录时告警(tmp_path, capsys):
    """把同一份数据同时当基线和对比，结果必然全为 0 —— 必须明确告警。

    否则用户看到一张全 0 的表，会误以为「没有任何劣化」，这是假结论。
    """
    result_dir = tmp_path / "结果"
    result_dir.mkdir()
    shutil.copy(str(FIXTURE / "基础性能测试" / "fio.txt"), str(result_dir / "fio.txt"))
    config_path = _写配置(tmp_path, result_dir, result_dir)

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    warnings = [line for line in output.splitlines() if line.startswith("    ⚠")]
    assert any("同一个目录" in line for line in warnings), "实际输出为：\n" + output


def test_命令行_两侧数据完全一致时告警(tmp_path, capsys):
    """两个不同目录、内容却一模一样的拷贝 —— 同样应当告警。"""
    for side in ("base", "copy"):
        target = tmp_path / side / "基础性能测试"
        target.mkdir(parents=True)
        shutil.copy(str(FIXTURE / "基础性能测试" / "fio.txt"), str(target / "fio.txt"))
    config_path = _写配置(tmp_path, tmp_path / "base", tmp_path / "copy")

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    warnings = [line for line in output.splitlines() if line.startswith("    ⚠")]
    assert any("完全一致" in line for line in warnings), "实际输出为：\n" + output


def _造两个结果目录(tmp_path, factor=0.8):
    """造一份 fio 对比：对比侧第 1 列乘 factor（<1 → 劣化）。

    目录结构照真实布局：`…/report/950/<OS>/64K/<侧别>`，
    这样两个路径的差异正好落在 OS 那一级。
    """
    baseline = tmp_path / "report" / "950" / "openEuler-24.03" / "64K" / "基线版本"
    candidate = tmp_path / "report" / "950" / "openEuler-24.03-SP1" / "64K" / "对比版本"
    for target in (baseline, candidate):
        target.mkdir(parents=True)
        shutil.copy(str(FIXTURE / "基础性能测试" / "fio.txt"), str(target / "fio.txt"))
    path = candidate / "fio.txt"
    lines = path.read_text(encoding="utf-8").splitlines()
    fields = lines[1].split()
    fields[1] = repr(float(fields[1]) * factor)
    lines[1] = " ".join(fields)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return baseline, candidate


def test_命令行_直接给两个目录无需配置文件(tmp_path, capsys):
    """需求方 2026-10-20：用传参代替配置文件，省掉每次改 JSON。"""
    baseline, candidate = _造两个结果目录(tmp_path)
    excel = tmp_path / "对比表格.xlsx"

    exit_code = main(
        [
            "--baseline",
            str(baseline),
            "--candidate",
            str(candidate),
            "--excel",
            str(excel),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "已写出对比表格" in output
    workbook = openpyxl.load_workbook(str(excel))
    header = [cell.value for cell in workbook["汇总"][1]][:3]
    # 名字按两侧路径差异推导：差异落在 OS 那一级 → openEuler-24.03-64K 等
    assert header == ["测试套", "openEuler-24.03-64K", "openEuler-24.03-SP1-64K"]


def test_命令行_传参指定显示名与内核版本(tmp_path, capsys):
    """OS 名称与内核版本可以直接用命令行给，优先于自动推导。"""
    baseline, candidate = _造两个结果目录(tmp_path)
    excel = tmp_path / "对比表格.xlsx"
    report = tmp_path / "报告"

    exit_code = main(
        [
            "--baseline",
            str(baseline),
            "--candidate",
            str(candidate),
            "--baseline-name",
            "openEuler-24.03-64K",
            "--candidate-name",
            "openEuler-24.03-SP1-64K",
            "--baseline-kernel",
            "6.8.0-1014-kunpeng",
            "--candidate-kernel",
            "6.8.0-1015-kunpeng",
            "--excel",
            str(excel),
            "--report",
            str(report),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    workbook = openpyxl.load_workbook(str(excel))
    assert [cell.value for cell in workbook["汇总"][1]][:3] == [
        "测试套",
        "openEuler-24.03-64K",
        "openEuler-24.03-SP1-64K",
    ]
    report_text = (tmp_path / "报告.txt").read_text(encoding="utf-8")
    assert "6.8.0-1014-kunpeng / 6.8.0-1015-kunpeng" in report_text
    assert "（配置）" not in output or "命令行" in output


def test_命令行_只给名字不给目录也能覆盖配置里的名字(tmp_path, capsys):
    """配置文件仍然可用，命令行名字覆盖它。"""
    baseline, candidate = _造两个结果目录(tmp_path)
    config_path = _写配置(tmp_path, baseline, candidate)
    excel = tmp_path / "对比表格.xlsx"

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--candidate-name",
            "命令行指定的名字",
            "--excel",
            str(excel),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    assert exit_code == 0
    workbook = openpyxl.load_workbook(str(excel))
    header = [cell.value for cell in workbook["汇总"][1]][:3]
    assert header[2] == "命令行指定的名字"


def test_命令行_既无配置也无目录时明确报错(tmp_path, capsys):
    exit_code = main(
        [
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    assert exit_code == 2
    err = capsys.readouterr().err
    assert "--baseline" in err and "--candidate" in err


def test_命令行_配置文件与目录同时给时明确报错(tmp_path, capsys):
    baseline, candidate = _造两个结果目录(tmp_path)
    config_path = _写配置(tmp_path, baseline, candidate)

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--baseline",
            str(baseline),
            "--candidate",
            str(candidate),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    assert exit_code == 2
    assert "只能二选一" in capsys.readouterr().err


def test_命令行_阈值可用传参覆盖(tmp_path, capsys):
    """传参 --threshold 覆盖配置里的阈值。"""
    baseline, candidate = _造两个结果目录(tmp_path, factor=0.90)   # -10%
    config_path = _写配置(tmp_path, baseline, candidate)             # 配置阈值 5%
    excel = tmp_path / "对比表格.xlsx"

    # 阈值提到 15% → 10% 的差异不再算劣化
    exit_code = main(
        [
            "--config",
            str(config_path),
            "--threshold",
            "15",
            "--excel",
            str(excel),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "劣化 0 条" in output


def _造单侧目录(tmp_path):
    target = tmp_path / "report" / "950" / "openEuler-24.03" / "64K" / "基线版本"
    target.mkdir(parents=True)
    for name in ("fio.txt", "stream.txt"):
        shutil.copy(str(FIXTURE / "基础性能测试" / name), str(target / name))
    return target


def test_命令行_只有基线时进入整理模式(tmp_path, capsys):
    """只有一份数据时仍可执行：只把结果整理成 excel，不做比对、不出报告。"""
    baseline = _造单侧目录(tmp_path)
    excel = tmp_path / "整理.xlsx"

    exit_code = main(
        [
            "--baseline",
            str(baseline),
            "--excel",
            str(excel),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "整理" in output
    assert excel.exists()

    workbook = openpyxl.load_workbook(str(excel))
    assert "子项劣化汇总" not in workbook.sheetnames
    assert len([cell.value for cell in workbook["汇总"][1]]) == 2      # 测试套 + 数值
    assert "fio" in workbook.sheetnames


def test_命令行_整理模式下请求报告会被跳过(tmp_path, capsys):
    """整理模式没有对比结果，报告无从写起 —— 明确说明并跳过，而不是报错。"""
    baseline = _造单侧目录(tmp_path)

    exit_code = main(
        [
            "--baseline",
            str(baseline),
            "--excel",
            str(tmp_path / "整理.xlsx"),
            "--report",
            str(tmp_path / "报告"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "跳过报告" in output or "不生成报告" in output
    assert not (tmp_path / "报告.txt").exists()
    assert not (tmp_path / "报告.html").exists()


def test_命令行_配置里没有候选时也走整理模式(tmp_path, capsys):
    baseline = _造单侧目录(tmp_path)
    config_path = tmp_path / "compare.json"
    config_path.write_text(
        json.dumps(
            {
                "baseline": {"path": str(baseline)},
                "candidates": [],
                "threshold_pct": 5.0,
            }
        ),
        encoding="utf-8",
    )

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(tmp_path / "整理.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    assert exit_code == 0
    workbook = openpyxl.load_workbook(str(tmp_path / "整理.xlsx"))
    assert "子项劣化汇总" not in workbook.sheetnames


def test_命令行_报告被排除的测试套(tmp_path, capsys):
    """只在一侧存在的套要明说被排除了，不能静默不出现。

    注意：断言必须**按行**匹配。pytest 会用函数名创建 tmp_path 目录，
    而 CLI 会把路径打印出来 —— 直接断言子串会被临时目录名意外满足（假通过）。
    """
    baseline = tmp_path / "baseline" / "基础性能测试"
    candidate = tmp_path / "candidate" / "基础性能测试"
    baseline.mkdir(parents=True)
    candidate.mkdir(parents=True)
    for name in ("fio.txt", "stream.txt"):
        shutil.copy(str(FIXTURE / "基础性能测试" / name), str(baseline / name))
    shutil.copy(str(FIXTURE / "基础性能测试" / "fio.txt"), str(candidate / "fio.txt"))
    config_path = _写配置(tmp_path, tmp_path / "baseline", tmp_path / "candidate")

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    skipped = [line for line in output.splitlines() if line.startswith("只在一侧出现")]
    assert skipped, "实际输出为：\n" + output
    assert "stream" in skipped[0]


def test_命令行_报告读到的套与未找到的套(tmp_path, capsys):
    """数据不全时要写清「读到哪些套、没找到哪些套」，便于定位是哪一侧缺数据。"""
    baseline = tmp_path / "baseline" / "基础性能测试"
    candidate = tmp_path / "candidate" / "基础性能测试"
    baseline.mkdir(parents=True)
    candidate.mkdir(parents=True)
    shutil.copy(str(FIXTURE / "基础性能测试" / "fio.txt"), str(baseline / "fio.txt"))
    shutil.copy(str(FIXTURE / "基础性能测试" / "fio.txt"), str(candidate / "fio.txt"))
    config_path = _写配置(tmp_path, tmp_path / "baseline", tmp_path / "candidate")

    exit_code = main(
        [
            "--config",
            str(config_path),
            "--excel",
            str(tmp_path / "对比表格.xlsx"),
            "--log-dir",
            str(tmp_path / "logs"),
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "fio" in output
    assert "未找到" in output
